"""Import trips from the Excel kilometerregistratie workbooks.

Each workbook has this layout on the ``Kilometer Registratie`` sheet:
  - B1: license plate (``Kenteken``), B2: year (``Periode``)
  - Row 6: header row
  - Row 7+: trip data

Columns (1-indexed):
  A=Datum, B=Rit nr., C=Beginstand KM, D=Eindstand KM, E=Van - Naar,
  F=Gereden Route, G=Bezoekadres, H=Karakter Rit (Z/P), I=Opmerkingen,
  (gap), K=Km Zakelijk, L=Privé Omgereden, M=Privé Kilometers

The parser tolerates messy Van-Naar formats (stray dots, trailing spaces,
missing separators) and skips incomplete rows (no Eindstand).
"""
from __future__ import annotations

import re
from dataclasses import dataclass, asdict, field
from datetime import datetime
from io import BytesIO
from typing import Optional

import openpyxl

from .storage import BUSINESS, PRIVATE


@dataclass
class ParsedTrip:
    date: str  # ISO-8601 date, e.g. "2026-04-24"
    km: float
    start_city: str
    end_city: str
    classification: str  # BUSINESS or PRIVATE
    license_plate: str
    note: Optional[str] = None
    start_odo: Optional[float] = None
    end_odo: Optional[float] = None
    raw_row: int = 0

    def to_dict(self) -> dict:
        return asdict(self)


@dataclass
class ImportResult:
    plate: str
    year: int
    trips: list[ParsedTrip] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    skipped: int = 0


def _split_van_naar(value: Optional[str]) -> tuple[str, str]:
    """Split 'City - City' into (start_city, end_city), handling stray dots and spaces."""
    if not value or not value.strip():
        return ("onbekend", "onbekend")

    cleaned = value.strip()
    # Remove stray dots before/after the separator (e.g. "Zwolle.- Enschede")
    cleaned = re.sub(r"\.\s*-\s*", " - ", cleaned)
    cleaned = re.sub(r"\s*-\.\s*", " - ", cleaned)

    parts = [p.strip() for p in cleaned.split("-", maxsplit=1)]
    if len(parts) == 2 and parts[0] and parts[1]:
        return (parts[0], parts[1])
    elif len(parts) == 1 and parts[0]:
        return (parts[0], parts[0])
    return ("onbekend", "onbekend")


def _parse_classification(value: Optional[str]) -> str:
    if isinstance(value, str) and value.strip().upper() == "Z":
        return BUSINESS
    return PRIVATE


def parse_workbook(file_bytes: bytes) -> ImportResult:
    """Parse an Excel kilometerregistratie workbook and return structured trip data."""
    wb = openpyxl.load_workbook(BytesIO(file_bytes), data_only=True)

    if "Kilometer Registratie" not in wb.sheetnames:
        return ImportResult(
            plate="",
            year=0,
            warnings=["Werkblad 'Kilometer Registratie' niet gevonden."],
        )

    ws = wb["Kilometer Registratie"]

    # Metadata
    plate = str(ws.cell(1, 2).value or "").strip()
    year_raw = ws.cell(2, 2).value
    try:
        year = int(float(year_raw)) if year_raw else datetime.now().year
    except (ValueError, TypeError):
        year = datetime.now().year

    result = ImportResult(plate=plate, year=year)

    for row_idx in range(7, ws.max_row + 1):
        datum = ws.cell(row_idx, 1).value
        if datum is None:
            continue  # empty row or formula-only row

        # Validate datum
        if isinstance(datum, datetime):
            date_str = datum.strftime("%Y-%m-%d")
        elif isinstance(datum, str):
            date_str = datum.strip()
        else:
            result.warnings.append(f"Rij {row_idx}: ongeldig datumformaat '{datum}', overgeslagen.")
            result.skipped += 1
            continue

        # Eindstand (column D) is required to compute km
        begin_raw = ws.cell(row_idx, 3).value
        end_raw = ws.cell(row_idx, 4).value

        if end_raw is None:
            result.warnings.append(
                f"Rij {row_idx}: geen eindstand, overgeslagen "
                f"({ws.cell(row_idx, 5).value or '?'})."
            )
            result.skipped += 1
            continue

        try:
            begin_odo = float(begin_raw) if begin_raw is not None else None
            end_odo = float(end_raw)
        except (ValueError, TypeError):
            result.warnings.append(f"Rij {row_idx}: ongeldige km-tellerstand, overgeslagen.")
            result.skipped += 1
            continue

        km = round(end_odo - begin_odo, 1) if begin_odo is not None else None
        if km is not None and km <= 0:
            result.warnings.append(
                f"Rij {row_idx}: negatieve of nul km ({km}), overgeslagen."
            )
            result.skipped += 1
            continue

        van_naar = ws.cell(row_idx, 5).value
        start_city, end_city = _split_van_naar(van_naar)
        classification = _parse_classification(ws.cell(row_idx, 8).value)
        note = ws.cell(row_idx, 9).value
        if isinstance(note, str):
            note = note.strip() or None

        trip = ParsedTrip(
            date=date_str,
            km=km if km is not None else 0.0,
            start_city=start_city,
            end_city=end_city,
            classification=classification,
            license_plate=plate,
            note=note,
            start_odo=begin_odo,
            end_odo=end_odo,
            raw_row=row_idx,
        )
        result.trips.append(trip)

    return result
